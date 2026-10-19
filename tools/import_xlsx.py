"""Rebuild js/data.js from the MS Basketball Drill & Practice Book spreadsheet.

Usage:  python3 tools/import_xlsx.py path/to/MS_Basketball_Drill_and_Practice_Book.xlsx

Reads the Drill Library, the four sample plans, the blueprint/principles on
Start Here, and the reference tabs. Diagrams live in js/diagrams-*.js and are
not touched; a new drill without a diagram shows its text only.
"""
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")  # openpyxl warns about the data-validation extension

OUT = Path(__file__).resolve().parent.parent / "js" / "data.js"

# Blueprint blocks from Start Here, grouped so the library reads like a practice.
GROUPS = [
    {"id": "warm", "name": "Warm-Up & Handles", "cats": ["WU", "BH"]},
    {"id": "skill", "name": "Pass, Finish & Shoot", "cats": ["PA", "FN", "SH"]},
    {"id": "defend", "name": "Defense & Boards", "cats": ["PD", "RB"]},
    {"id": "reads", "name": "Small-Sided & Post", "cats": ["SS", "PO"]},
    {"id": "team", "name": "Team & Transition", "cats": ["TR", "TO", "TD", "SP"]},
    {"id": "compete", "name": "Compete & Condition", "cats": ["CG", "CD"]},
    {"id": "breaks", "name": "Breaks & Wrap-Up", "cats": ["BR"]},
]


def text(v):
    return "" if v is None else str(v).strip()


def rows(ws, **kw):
    for r in ws.iter_rows(values_only=True, **kw):
        yield [text(c) if not (isinstance(c, str) and c.startswith("=")) else "" for c in r]


def drill_ids(s):
    """'PD-09, TD-03' -> ids; 'BH-01 to BH-08' -> the full range."""
    m = re.match(r"([A-Z]{2})-(\d+) to [A-Z]{2}-(\d+)", s)
    if m:
        return [f"{m[1]}-{i:02d}" for i in range(int(m[2]), int(m[3]) + 1)]
    return re.findall(r"[A-Z]{2}-\d{2}", s)


def main(path):
    wb = openpyxl.load_workbook(path)

    # --- Start Here: categories, blueprint, principles, source links ---
    cats, blueprint, principles, sources = {}, [], [], []
    section = None
    for r in rows(wb["Start Here"]):
        b, c, d = r[1], r[2], r[3]
        if b and b.isupper() and not c:
            section = b
            continue
        if section == "CATEGORIES" and b:
            code, name = b.split(None, 1)
            cats[code] = name.strip()
        elif section == "90-MINUTE PRACTICE BLUEPRINT" and b and b != "Block":
            blueprint.append({"block": b, "min": c, "cats": re.findall(r"[A-Z]{2}(?:-\d{2})?", d)})
        elif section == "PRINCIPLES FOR PUSHING MIDDLE SCHOOLERS" and b:
            principles.append(b.lstrip("• ").strip())
        elif section == "SOURCES FOR THE COLLEGE DRILLS" and d.startswith("http"):
            sources.append((b, d))

    # --- Drill Library ---
    keys = ["id", "cat", "name", "min", "level", "college", "players", "setup", "runs", "points", "ms", "source"]
    code_by_name = {v: k for k, v in cats.items()}
    drills = []
    for r in wb["Drill Library"].iter_rows(min_row=5, values_only=True):
        if not r[0] or not r[2]:
            continue
        d = {k: text(v) for k, v in zip(keys, r[:12])}
        d["cat"] = d["id"].split("-")[0] if "-" in d["id"] else code_by_name.get(d["cat"], d["cat"])
        d["min"] = int(float(d["min"] or 5))
        d["points"] = [p.lstrip("•- ").strip() for p in d["points"].split("\n") if p.strip()]
        # A source row like "Shell Progression & Three Stops" covers several drills: match each part by its words.
        words = set(re.findall(r"[a-z0-9]+", d["name"].lower()))
        for src_name, url in sources:
            parts = [re.findall(r"[a-z0-9]+", p.lower()) for p in re.split(r" / |, | & ", re.sub(r"\(.*?\)", "", src_name))]
            if any(len(p) >= 2 and sum(w in words for w in p) / len(p) >= 0.66 for p in parts):
                d["url"] = url
                break
        drills.append({k: v for k, v in d.items() if v not in ("", None, [])})
    by_name = {d["name"]: d["id"] for d in drills}

    # --- Sample plans ---
    templates = []
    for ws in wb.worksheets:
        if not ws.title.startswith("Plan "):
            continue
        all_rows = list(rows(ws, max_row=36))
        title = all_rows[0][0].split(":", 1)[-1].strip().title()
        focus = next((r[1] for r in all_rows if r[0] == "Focus"), "")
        message = next((r[1] for r in all_rows if r[0] == "Message"), "")
        items = [by_name[r[2]] for r in all_rows[10:30] if r[2] in by_name]
        templates.append({"id": ws.title.split(" - ")[0].lower().replace(" ", ""), "name": ws.title.split(" - ", 1)[1],
                          "title": title, "focus": focus, "message": message, "items": items})

    # --- Play Refinement: section rows have a title in A and 'Rep It With' further right ---
    plays = []
    for r in rows(wb["Play Refinement"], min_row=4):
        if "Rep It With" in r:
            plays.append({"title": r[0].title(), "items": []})
        elif r[0] and not r[1] and plays:
            plays[-1]["items"].append({"sub": r[0]})
        elif r[1] and plays:
            plays[-1]["items"].append({"n": r[0], "title": r[1], "text": r[2], "drills": drill_ids(r[3])})

    # --- Defense Concepts ---
    defense = []
    for r in rows(wb["Defense Concepts"], min_row=5):
        if r[0] and not r[1]:
            defense.append({"title": r[0].title().replace("(New)", "").strip(), "items": []})
        elif r[1] and defense:
            defense[-1]["items"].append({"title": r[1], "staff": r[2], "text": r[3], "drills": drill_ids(r[4])})

    # --- Position Skills ---
    positions = []
    for r in rows(wb["Position Skills"], min_row=4):
        if r[1] == "Staff Description (original)":
            positions.append({"title": r[0].title(), "items": []})
        elif r[0] and positions:
            positions[-1]["items"].append({"title": r[0], "staff": r[1], "text": r[2], "drills": drill_ids(r[3])})

    # --- HS Readiness benchmarks ---
    tests = []
    for r in rows(wb["HS Readiness"], min_row=5, max_row=13):
        if r[0]:
            tests.append({"test": r[0], "how": r[1], "target": r[4], "better": r[5], "drills": drill_ids(r[6])})

    data = {
        "groups": GROUPS,
        "cats": cats,
        "blueprint": blueprint,
        "principles": principles,
        "templates": templates,
        "plays": plays,
        "defense": defense,
        "positions": positions,
        "tests": tests,
    }
    lines = ["// Generated by tools/import_xlsx.py from the Drill & Practice Book. Re-run it after editing the spreadsheet.",
             "const DATA = " + json.dumps(data, ensure_ascii=False) + ";",
             "DATA.drills = ["]
    lines += ["  " + json.dumps(d, ensure_ascii=False) + "," for d in drills]
    lines.append("];")
    OUT.write_text("\n".join(lines) + "\n")
    print(f"Wrote {OUT}: {len(drills)} drills, {len(templates)} plans, {sum(1 for d in drills if 'url' in d)} source links")


if __name__ == "__main__":
    main(sys.argv[1])
